"""C1158 deck: one 30-min lunch after Activity 2 (12:10 in the LP); drop End of Day 1 and the 2nd lunch."""
from pptx import Presentation
from pptx.util import Emu

DECK = "courseware/Improve Employee Wellness and Work Resilience-v1.0.pptx"
prs = Presentation(DECK)
S = prs.slides

def txt(s): return [sh.text_frame.text for sh in s.shapes if sh.has_text_frame]

def set_one(slide, sid, text):
    sh = next(s for s in slide.shapes if s.shape_id == sid)
    p = sh.text_frame.paragraphs[0]
    p.runs[0].text = text
    for r in p.runs[1:]: r.text = ""

assert "Lunch Break" in txt(S[43]), txt(S[43])
set_one(S[43], 7, "30 minutes")

drop = [i for i, s in enumerate(S) if "End of Day 1" in txt(s) or ("Lunch Break" in txt(s) and "60 minutes" in txt(s))]
assert drop == [78, 87], drop
lst = prs.slides._sldIdLst
for i in sorted(drop, reverse=True):
    sid = lst[i]
    prs.part.drop_rel(sid.rId)
    lst.remove(sid)

# Renumber the static bottom-right page-number footers.
W, H = prs.slide_width, prs.slide_height
for n, s in enumerate(prs.slides, start=1):
    for sh in s.shapes:
        if sh.has_text_frame and sh.text_frame.text.strip().isdigit() \
           and sh.top > H * 0.85 and sh.left > W * 0.5:
            p = sh.text_frame.paragraphs[0]
            p.runs[0].text = str(n)
            for r in p.runs[1:]: r.text = ""
prs.save(DECK)
print("slides:", len(prs.slides))
