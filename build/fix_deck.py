"""Post-convert fixes for the C1158 deck (run after convert_deck.py)."""
import re
import sys
from pptx import Presentation

DECK = "courseware/Improve Employee Wellness and Work Resilience-v1.0.pptx"
OLD_URL = "https://www.tertiarycourses.com.sg/wsq-empowering-employee-health-and-wellness-at-the-workplace.html"
NEW_URL = "https://www.tertiarycourses.com.sg/improve-employee-wellness-and-work-resilience.html"
OLD_REPO = "https://github.com/tertiarycourses/TGS-2024045222-Empowering-Employee-Health-and-Wellness-at-the-Workplace"
NEW_REPO = "https://github.com/tertiarycourses/C1158-Improve-Employee-Wellness-and-Work-Resilience"
KA = re.compile(r"\s*·\s*[KA]\d\b")


def set_para(para, text):
    runs = para.runs
    runs[0].text = text
    for r in runs[1:]:
        r.text = ""


def set_shape(slide, sid, lines):
    sh = next(s for s in slide.shapes if s.shape_id == sid)
    paras = sh.text_frame.paragraphs
    assert len(paras) == len(lines), (sid, len(paras), lines)
    for p, t in zip(paras, lines):
        set_para(p, t)


prs = Presentation(DECK)
S = prs.slides

# Global: strip K#/A# competency codes (keep LO#), rename TSC-derived eyebrow.
for slide in S:
    for sh in slide.shapes:
        if not sh.has_text_frame:
            continue
        for p in sh.text_frame.paragraphs:
            t = "".join(r.text for r in p.runs)
            if not p.runs:
                continue
            new = KA.sub("", t)
            if new == "C1158 Non-WSQ commercial short course":
                new = "KEY CONCEPTS"
            if new != t:
                set_para(p, new)
    if slide.has_notes_slide:
        for p in slide.notes_slide.notes_text_frame.paragraphs:
            t = "".join(r.text for r in p.runs)
            if OLD_URL in t:
                set_para(p, t.replace(OLD_URL, NEW_URL))

# Cover date
set_shape(S[0], 14, ["Version v1.0  ·  2 October 2026"])
# Course provider card
set_shape(S[2], 12, ["Training provider"])
set_shape(S[2], 27, ["PROGRAMME", "Professional short course"])
# Course facts
set_shape(S[3], 21, ["DELIVERY", "Professional development and workplace application"])
# Ground rules
set_shape(S[5], 36, ["Attendance", "Attend the full day to get the most from the course."])
# Outline: 1 day, Morning / Afternoon
set_shape(S[6], 9, ["Morning"])
set_shape(S[6], 10, ["Afternoon"])
set_shape(S[6], 11, ["•  Morning — Understand the evidence",
                     "–  Topic 1: requirements, factors and wellbeing dimensions",
                     "–  Topic 1: trends, technology and evidence",
                     "–  Activity 1: privacy-safe wearable pilot",
                     "–  Activity 2: segmented workforce needs"])
set_shape(S[6], 12, ["•  Afternoon — Design, implement and improve",
                     "–  Activity 3, then Topic 2: programme design and implementation",
                     "–  Activities 4–7: resilience, portfolio, playbook, participation",
                     "–  Topic 3: communication, feedback and evaluation",
                     "–  Activities 8–10 and course wrap-up"])
# Source map card
src = S[127]
sh = next(s for s in src.shapes if s.shape_id == 36)
assert sh.text_frame.text.startswith("Course page"), sh.text_frame.text
set_shape(src, 36, ["Course page", "Course outcomes, topics and duration."])
# Course Feedback (was TRAQOM digital attendance)
fb = S[129]
set_shape(fb, 4, ["COURSE FEEDBACK"])
set_shape(fb, 5, ["Course Feedback"])
set_shape(fb, 11, ["Scan", "Scan the course feedback QR code shown by the trainer."])
set_shape(fb, 16, ["Rate", "Rate the content, activities and trainer delivery."])
set_shape(fb, 21, ["Suggest", "Tell us what would make the course more useful at work."])
set_shape(fb, 26, ["Certificate", "Your certificate of completion is emailed after the course."])

# Hyperlink targets (python-pptx caches target_ref, so set both).
for slide in S:
    for rel in slide.part.rels.values():
        if rel.is_external and rel.target_ref.startswith(OLD_REPO):
            new = rel.target_ref.replace(OLD_REPO, NEW_REPO)
            rel._target = new
            rel.__dict__["target_ref"] = new

prs.save(DECK)
print("fixed", DECK, len(S), "slides")
