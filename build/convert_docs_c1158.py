"""In-place WSQ -> non-WSQ conversion of the LP and LG for C1158.

Source: source-wsq/courseware/{LP,LG}-Empowering Employee Health and Wellness at the Workplace.docx
Output: courseware/{LP,LG}-Improve Employee Wellness and Work Resilience.docx

The WSQ course is 2 days; C1158 is ONE day (9:30am-5:30pm, 450 min + 30-min lunch).
Every WSQ topic/activity row is kept and rescaled proportionally (see SCHEDULE).
"""
import copy
import re
from docx import Document

OLD_TITLE = "Empowering Employee Health and Wellness at the Workplace"
NEW_TITLE = "Improve Employee Wellness and Work Resilience"
OLD_URL = "https://www.tertiarycourses.com.sg/wsq-empowering-employee-health-and-wellness-at-the-workplace.html"
NEW_URL = "https://www.tertiarycourses.com.sg/improve-employee-wellness-and-work-resilience.html"
VERSION_ROW = ["1.0", "2 October 2026",
               "One-day course courseware with ten real-use-case activity packs and detailed learner procedures.",
               "Dr. Alfred Ang"]

# Source deck slides (0-based) removed from the converted deck.
# (src 50, the first lunch slide, is kept as the 30-min lunch; 85 = End of Day 1, 94 = 2nd lunch.)
DROPPED = {2, 8, 9, 10, 11, 12, 14, 85, 94, 136, 137, 138}


def slide_no(n):
    """Map a 1-based source slide number to the converted deck."""
    i = n - 1
    assert i not in DROPPED, n
    return n - sum(1 for d in DROPPED if d < i)


def set_text(container, text):
    paras = container.paragraphs
    p = paras[0]
    if p.runs:
        p.runs[0].text = text
        for r in p.runs[1:]:
            r.text = ""
    else:
        p.add_run(text)
    for extra in paras[1:]:
        extra._element.getparent().remove(extra._element)


def drop(el):
    el.getparent().remove(el)


def common(doc):
    for p in doc.paragraphs:
        t = p.text
        if t == OLD_TITLE:
            set_text_para(p, NEW_TITLE)
        elif t == "TGS Ref No: TGS-2024045222":
            set_text_para(p, "Course Code: C1158")
        elif t == "Version 8.0":
            set_text_para(p, "Version 1.0")
        elif OLD_TITLE in t:
            set_text_para(p, t.replace(OLD_TITLE, NEW_TITLE))
    # Hyperlinks: python-docx's p.runs skips <w:hyperlink> runs, so rewrite by rel target.
    W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
    R = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"
    rels = doc.part.rels
    for h in doc.element.body.iter(W + "hyperlink"):
        rid = h.get(R + "id")
        if rid and rels[rid].is_external and rels[rid].target_ref == OLD_URL:
            rels[rid]._target = NEW_URL
            for t in h.iter(W + "t"):
                t.text = t.text.replace(OLD_URL, NEW_URL)
    vt = doc.tables[0]
    assert vt.rows[0].cells[0].text == "Version Number"
    for c, v in zip(vt.rows[1].cells, VERSION_ROW):
        set_text(c, v)
    for r in list(vt.rows)[2:]:
        drop(r._tr)


def set_text_para(p, text):
    if p.runs:
        p.runs[0].text = text
        for r in p.runs[1:]:
            r.text = ""
    else:
        p.add_run(text)


def delete_section(doc, heading, stop_style="Heading 1"):
    """Delete a Heading-1 section (heading + body) up to the next Heading 1."""
    body = doc.element.body
    start = next(p for p in doc.paragraphs if p.text == heading and p.style.name == stop_style)
    el = start._element
    nxt = el.getnext()
    drop(el)
    while nxt is not None:
        if nxt.tag.endswith("}p"):
            sty = nxt.find(".//{http://schemas.openxmlformats.org/wordprocessingml/2006/main}pStyle")
            if sty is not None and sty.get("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}val") in ("Heading1",):
                break
        if nxt.tag.endswith("}sectPr"):
            break
        cur, nxt = nxt, nxt.getnext()
        drop(cur)


MAPPING = re.compile(r"^(?:TSC mapping|Mapping): ([^.]*)\.\s*(.*)$")


def mapping_line(t):
    m = MAPPING.match(t)
    if not m:
        return None
    codes, rest = m.groups()
    los = [c for c in re.split(r"\s*·\s*", codes) if c.startswith("LO")]
    if los:
        return f"Outcome: {' · '.join(los)}. {rest}"
    return f"Focus: {rest}"


# ------------------------------------------------------------------ LG ------
lg = Document(f"source-wsq/courseware/LG-{OLD_TITLE}.docx")
common(lg)
for p in lg.paragraphs:
    new = mapping_line(p.text)
    if new:
        set_text_para(p, new)
    if p.text == "Course Learning Outcomes and TSC Mapping":
        set_text_para(p, "Course Learning Outcomes")
# K/A competency lists (Knowledge / Abilities sub-sections under the LO heading)
paras = lg.paragraphs
k = next(i for i, p in enumerate(paras) if p.text == "Knowledge" and p.style.name == "Heading 2")
end = next(i for i, p in enumerate(paras) if i > k and p.style.name == "Heading 1")
for p in paras[k:end]:
    drop(p._element)
delete_section(lg, "Assessment Preparation")
lg.save(f"courseware/LG-{NEW_TITLE}.docx")

# ------------------------------------------------------------------ LP ------
lp = Document(f"source-wsq/courseware/LP-{OLD_TITLE}.docx")
common(lp)
OVERVIEW_OLD = "This two-day WSQ course develops"
for p in lp.paragraphs:
    t = p.text
    if t.startswith(OVERVIEW_OLD):
        set_text_para(p, t.replace(OVERVIEW_OLD, "This one-day course develops"))
    elif t == "Day 1 — Evidence, Needs and Programme Design":
        set_text_para(p, "Day 1 · 9:30am – 5:30pm")
    elif t.startswith("Total training/assessment time excluding lunch"):
        set_text_para(p, "Day total: 7.5 instructional hours (9:30am – 5:30pm), with one 30-minute lunch.")
    elif t == "LMS access for courseware and assessment":
        set_text_para(p, "LMS access for courseware")
    else:
        new = mapping_line(t)
        if new:
            set_text_para(p, new)
        m = re.match(r"^(Topic \d — .*) · Slides (\d+)–(\d+)$", t)
        if m:
            set_text_para(p, f"{m.group(1)} · Slides {slide_no(int(m.group(2)))}–{slide_no(int(m.group(3)))}")
        m = re.match(r"^(Activity \d+ — .*) · Slide (\d+)$", t)
        if m:
            set_text_para(p, f"{m.group(1)} · Slide {slide_no(int(m.group(2)))}")

# Day 2 heading + its total line + its table
paras = lp.paragraphs
d2 = next(i for i, p in enumerate(paras) if p.text.startswith("Day 2 — "))
d2_total = paras[d2 + 1]
assert d2_total.text.startswith("Total training/assessment time") or d2_total.text.startswith("Day total"), d2_total.text
drop(paras[d2]._element)
drop(d2_total._element)
drop(lp.tables[3]._tbl)
delete_section(lp, "Assessment")

# Course facts table
facts = lp.tables[1]
for r in list(facts.rows):
    key = r.cells[0].text
    if key in ("TSC", "Assessment"):
        drop(r._tr)
    elif key == "Course Code":
        set_text(r.cells[1], "C1158")
    elif key == "Duration":
        set_text(r.cells[1], "1 day · 7.5 hours")
    elif key == "Training":
        set_text(r.cells[1], "7.5 hours")
    elif key == "Daily Time":
        set_text(r.cells[1], "9:30am – 5:30pm")

# One-day schedule: WSQ teaching rows rescaled onto 450 min, one 30-min lunch.
# (source minutes) Admin 30, T1a 75, A1-2 85, T1b 90 | A3-5 130, T2 70, A6-7 75,
# T3 70, A8-9 50, A10 60, synthesis 50  -> sum 785, scaled x450/785 at 5-min steps.
PRES = "Interactive presentation · case study · discussion"
SCHEDULE = [
    ("welcome", 15, "Welcome, introductions, course outline and ground rules", "Interactive presentation", "—"),
    ("topic", 45, "Topic 1 — health, wellness, evidence and requirements", PRES, f"Slides {slide_no(20)}–{slide_no(56)}"),
    ("topic", 50, "Activities 1–2 — privacy and workforce needs", PRES, f"A1: {slide_no(43)}, A2: {slide_no(47)}"),
    ("topic", 50, "Topic 1 — factors, trends, technology and research", PRES, f"Slides {slide_no(20)}–{slide_no(56)}"),
    ("lunch", 30, "Lunch", "Lunch", "—"),
    ("topic", 70, "Activities 3–5 — technology, resilience and programme portfolio", PRES,
     f"A3: {slide_no(52)}, A4: {slide_no(78)}, A5: {slide_no(82)}"),
    ("topic", 40, "Topic 2 — programme design and implementation", PRES, f"Slides {slide_no(57)}–{slide_no(96)}"),
    ("topic", 45, "Activities 6–7 — policy-to-practice and participation", PRES, f"A6: {slide_no(87)}, A7: {slide_no(91)}"),
    ("topic", 40, "Topic 3 — communication, feedback and evaluation", PRES, f"Slides {slide_no(97)}–{slide_no(141)}"),
    ("topic", 30, "Activities 8–9 — campaign and feedback system", PRES, f"A8: {slide_no(118)}, A9: {slide_no(122)}"),
    ("topic", 35, "Activity 10 — management review and integrated improvement", PRES, f"A10: {slide_no(126)}"),
    ("wrap", 30, "Course synthesis, evidence review and course feedback", "Reflection · Q&A", "—"),
]
assert sum(m for k, m, *_ in SCHEDULE if k != "lunch") == 450
assert sum(m for _, m, *_ in SCHEDULE) == 480

sched = lp.tables[2]
rows = list(sched.rows)
assert rows[0].cells[0].text == "Time"
tmpl = {"welcome": rows[1]._tr, "topic": rows[2]._tr, "lunch": rows[5]._tr, "wrap": rows[9]._tr}
assert rows[5].cells[2].text == "Lunch" and rows[9].cells[2].text.startswith("Day 1 synthesis")
tmpl = {k: copy.deepcopy(v) for k, v in tmpl.items()}
for r in rows[1:]:
    drop(r._tr)

t = 9 * 60 + 30
for kind, mins, what, method, slides in SCHEDULE:
    tr = copy.deepcopy(tmpl[kind])
    sched._tbl.append(tr)
    row = sched.rows[-1]
    a, b = t, t + mins
    t = b
    vals = [f"{a // 60:02d}:{a % 60:02d}–{b // 60:02d}:{b % 60:02d}", str(mins), what, method, slides]
    for c, v in zip(row.cells, vals):
        set_text(c, v)
assert t == 17 * 60 + 30, t
lp.save(f"courseware/LP-{NEW_TITLE}.docx")
print("ok")
